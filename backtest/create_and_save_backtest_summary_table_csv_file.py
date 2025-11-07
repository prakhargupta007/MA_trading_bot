'''
# BEFORE AUTOMATED BUNCH BACKTESTING EDITS
import csv 
from tabulate import tabulate 
import pandas as pd
import os

from config import TICKERS
from config import FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE
FOLDER_PATH = FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE

def create_and_save_backtest_summary_table_csv_file(
    summary_tickers, 
    summary_CAGRs, 
    summary_buy_hold_CAGRs, 
    summary_cagr_strategy_efficiencies, 
    summary_sharpe_ratios, 
    summary_realized_CAGRs=None,
    summary_realized_sharpe_ratios=None,
    summary_realized_cagr_strategy_efficiencies=None
):
    os.makedirs(FOLDER_PATH, exist_ok=True)

    file_name = f'summary_backtest_table_{TICKERS}.csv'
    full_path = os.path.join(FOLDER_PATH, file_name)

    # Write to CSV file
    with open(full_path, mode='w', newline='') as file:
        writer = csv.writer(file)
        writer.writerow([
            'Ticker', 
            'CAGR', 
            'Realized CAGR', 
            'B&H CAGR', 
            'CAGR Eff.',
            'Realized CAGR Eff.',
            'Sharpe', 
            'Realized Sharpe'
        ])

        for ticker, strategy_cagr, realized_cagr, buy_hold_cagr, cagr_strategy_efficiency, realized_cagr_efficiency, sharpe_ratio, realized_sharpe_ratio in zip(
            summary_tickers, 
            summary_CAGRs, 
            summary_realized_CAGRs or [None] * len(summary_tickers),
            summary_buy_hold_CAGRs, 
            summary_cagr_strategy_efficiencies, 
            summary_realized_cagr_strategy_efficiencies or [None] * len(summary_tickers),
            summary_sharpe_ratios, 
            summary_realized_sharpe_ratios or [None] * len(summary_tickers)
        ):
            writer.writerow([
                ticker, 
                f"{strategy_cagr:.2f}" if strategy_cagr is not None else "N/A", 
                f"{realized_cagr:.2f}" if realized_cagr is not None else "N/A", 
                f"{buy_hold_cagr:.2f}" if buy_hold_cagr is not None else "N/A", 
                f"{cagr_strategy_efficiency:.2f}" if cagr_strategy_efficiency is not None else "N/A", 
                f"{realized_cagr_efficiency:.2f}" if realized_cagr_efficiency is not None else "N/A", 
                f"{sharpe_ratio:.2f}" if sharpe_ratio is not None else "N/A", 
                f"{realized_sharpe_ratio:.2f}" if realized_sharpe_ratio is not None else "N/A"
            ])

    print(f'\n\033[1mSummary of all tickers:\033[0m')
    df = pd.read_csv(full_path)
    summary_table = tabulate(df, headers='keys', tablefmt='grid', floatfmt=".2f")

    return summary_table, full_path
'''




'''
# AFTER AUTOMATED BUNCH BACKTESTING EDITS
import csv
import os
import pandas as pd
import numpy as np
from tabulate import tabulate

from config import TICKERS
from config import FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE
FOLDER_PATH = FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE


# ---------------- Helper Metric Functions ---------------- #

def max_drawdown(equity_curve: pd.Series) -> float:
    if equity_curve is None or equity_curve.empty:
        return 0.0
    roll_max = equity_curve.cummax()
    dd = equity_curve / roll_max - 1
    return float(100 * dd.min())  # in percent


def sortino_ratio(equity_curve: pd.Series, risk_free_rate_annual: float = 0.0, periods_per_year: int = 252) -> float:
    if equity_curve is None or equity_curve.empty:
        return 0.0
    returns = equity_curve.pct_change().dropna()
    if len(returns) < 2:
        return 0.0
    rf = risk_free_rate_annual / periods_per_year
    excess = returns - rf
    downside = excess[excess < 0]
    if downside.std(ddof=1) == 0:
        return 0.0
    return float((excess.mean() / downside.std(ddof=1)) * np.sqrt(periods_per_year))


def calmar_ratio(cagr_percent: float, max_dd_percent: float) -> float:
    if max_dd_percent == 0:
        return 0.0
    return float(cagr_percent / abs(max_dd_percent))


def trade_stats(trades_df: pd.DataFrame):
    """Return win_rate(%), avg_win, avg_loss, profit_factor, num_trades."""
    if trades_df is None or trades_df.empty or "Profit" not in trades_df.columns:
        return 0.0, 0.0, 0.0, 0.0, 0
    wins = trades_df[trades_df["Profit"] > 0]["Profit"]
    losses = trades_df[trades_df["Profit"] < 0]["Profit"]

    win_rate = 100 * len(wins) / max(1, len(trades_df))
    avg_win = float(wins.mean()) if len(wins) else 0.0
    avg_loss = float(losses.mean()) if len(losses) else 0.0
    gross_profit = wins.sum()
    gross_loss = abs(losses.sum())
    profit_factor = float(gross_profit / gross_loss) if gross_loss > 0 else 0.0
    num_trades = len(trades_df)

    return win_rate, avg_win, avg_loss, profit_factor, num_trades


# ---------------- Main Summary Writer ---------------- #

def create_and_save_backtest_summary_table_csv_file(
    summary_tickers,
    summary_CAGRs,
    summary_buy_hold_CAGRs,
    summary_cagr_strategy_efficiencies,
    summary_sharpe_ratios,
    summary_buy_hold_sharpes,
    summary_sharpe_efficiencies,
    portfolio_values_dict=None,
    buy_hold_equity_dict=None,
    trade_tables_dict=None,
    ml_metrics_dict=None,  # NEW: {ticker: {acc, prec, rec, f1, conf_matrix}}
    strategy_names=None,
    model_numbers=None,
    feature_set_ids=None,
    plot_links=None
):
    """
    Saves a full-featured summary CSV including all trading & ML metrics.
    """
    os.makedirs(FOLDER_PATH, exist_ok=True)
    file_name = f"summary_backtest_table_{TICKERS}.csv"
    full_path = os.path.join(FOLDER_PATH, file_name)

    with open(full_path, mode="w", newline="") as file:
        writer = csv.writer(file)

        writer.writerow([
            "Ticker",
            "Strategy",
            "Model_Number",
            "Feature_Set_ID",
            "CAGR(%)",
            "B&H CAGR(%)",
            "CAGR Efficiency",
            "Sharpe",
            "B&H Sharpe",
            "Sharpe Efficiency",
            "Sortino",
            "Max_Drawdown(%)",
            "Calmar",
            "WinRate(%)",
            "Avg_Win",
            "Avg_Loss",
            "Profit_Factor",
            "Number_of_Trades",
            "Accuracy",
            "Precision",
            "Recall",
            "F1_Score",
            "Confusion_Matrix",
            "Plot_Link"











        ])

        for i, ticker in enumerate(summary_tickers):
            cagr = summary_CAGRs[i]
            bh_cagr = summary_buy_hold_CAGRs[i]
            cagr_eff = summary_cagr_strategy_efficiencies[i]
            sharpe = summary_sharpe_ratios[i]
            bh_sharpe = summary_buy_hold_sharpes[i]
            sharpe_eff = summary_sharpe_efficiencies[i]

            strategy = strategy_names[i] if strategy_names else ""
            model_no = model_numbers[i] if model_numbers else ""
            feat_id = feature_set_ids[i] if feature_set_ids else ""
            plot_link = plot_links[i] if plot_links else ""

            # --- trading metrics ---
            pv_series = portfolio_values_dict.get(ticker) if portfolio_values_dict else None
            bh_equity = buy_hold_equity_dict.get(ticker) if buy_hold_equity_dict else None
            trades_df = trade_tables_dict.get(ticker) if trade_tables_dict else None

            winrate, avg_win, avg_loss, profit_factor, num_trades = trade_stats(trades_df)
            mdd = max_drawdown(pv_series)
            sortino = sortino_ratio(pv_series)
            calmar = calmar_ratio(cagr, mdd)

            # --- ML metrics (if available) ---
            ml_metrics = ml_metrics_dict.get(ticker) if ml_metrics_dict else {}
            acc = ml_metrics.get("accuracy", "")
            prec = ml_metrics.get("precision", "")
            rec = ml_metrics.get("recall", "")
            f1 = ml_metrics.get("f1", "")
            conf_matrix = ml_metrics.get("confusion_matrix", "")

            writer.writerow([
                ticker,
                strategy,
                model_no,
                feat_id,
                f"{cagr:.2f}",
                f"{bh_cagr:.2f}",
                f"{cagr_eff:.2f}",
                f"{sharpe:.2f}",
                f"{bh_sharpe:.2f}",
                f"{sharpe_eff:.2f}",
                f"{sortino:.2f}",
                f"{mdd:.2f}",
                f"{calmar:.2f}",
                f"{winrate:.2f}",
                f"{avg_win:.2f}",
                f"{avg_loss:.2f}",
                f"{profit_factor:.2f}",
                num_trades,
                acc,
                prec,
                rec,
                f1,
                conf_matrix,
                plot_link
            ])

    print(f"\n\033[1mSummary of all tickers:\033[0m")
    df = pd.read_csv(full_path)
    summary_table = tabulate(df, headers="keys", tablefmt="grid", floatfmt=".2f")
    return summary_table, full_path
'''

















'''
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font
from config import TICKERS

FINAL_EXCEL_DIR = "/Users/prakhar/Desktop/MA_trading_bot/automated_backtesting_results/full_backtesting_tables_excel"

def create_and_save_backtest_summary_table_excel(all_results):
    """
    Saves all backtesting and ML metrics into one Excel file.
    all_results is a list of dictionaries, each containing:
        {
            "Ticker": str,
            "Strategy": str,
            "Model": str or None,
            "Feature_Set_ID": str,
            "CAGR": float,
            "BH_CAGR": float,
            "CAGR_Eff": float,
            "Sharpe": float,
            "BH_Sharpe": float,
            "Sharpe_Eff": float,
            "Sortino": float,
            "Max_Drawdown": float,
            "Calmar": float,
            "WinRate": float,
            "Profit_Factor": float,
            "Num_Trades": int,
            "Accuracy": float or None,
            "Precision": float or None,
            "Recall": float or None,
            "F1": float or None,
            "Confusion_Matrix": str or None,
            "Plot_Link": str or None
        }
    """

    os.makedirs(FINAL_EXCEL_DIR, exist_ok=True)
    file_name = f"final_backtesting_summary_{TICKERS}.xlsx"
    full_path = os.path.join(FINAL_EXCEL_DIR, file_name)

    df = pd.DataFrame(all_results)

    # Convert HTML paths into clickable hyperlinks for Excel
    def make_hyperlink(path):
        if pd.isna(path) or not path:
            return ""
        return f'=HYPERLINK("{path}", "Open Plot")'

    df["Plot_Link"] = df["Plot_Link"].apply(make_hyperlink)

    # Save Excel
    with pd.ExcelWriter(full_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Summary")

        # Format header row bold
        ws = writer.sheets["Summary"]
        for cell in ws[1]:
            cell.font = Font(bold=True)

    print(f"\n✅ Final backtesting summary saved at:\n{full_path}\n")
    return full_path
''' 













import os
import math
import re
import pandas as pd
import numpy as np
from tabulate import tabulate
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import (
    TICKERS,
    RISK_FREE_RATE,
    STARTING_BALANCE,
    FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE,
)

# Where to save the final Excel
FINAL_EXCEL_DIR = "/Users/prakhar/Desktop/MA_trading_bot/automated_backtesting_results/full_backtesting_tables_excel"

# ---------------- Helper Metric Functions (local) ---------------- #

def _max_drawdown(equity_curve: pd.Series) -> float:
    if equity_curve is None or len(equity_curve) == 0:
        return 0.0
    eq = pd.Series(equity_curve).astype(float).reset_index(drop=True)
    roll_max = eq.cummax()
    dd = eq / roll_max - 1.0
    return float(100.0 * dd.min())  # percent

def _sortino_ratio(equity_curve: pd.Series, rf_annual: float = 0.0, periods_per_year: int = 252) -> float:
    if equity_curve is None or len(equity_curve) < 3:
        return 0.0
    eq = pd.Series(equity_curve).astype(float).reset_index(drop=True)
    rets = eq.pct_change().dropna()
    if len(rets) < 2:
        return 0.0
    rf = rf_annual / periods_per_year
    excess = rets - rf
    downside = excess[excess < 0]
    denom = downside.std(ddof=1)
    if denom == 0 or np.isnan(denom):
        return 0.0
    return float((excess.mean() / denom) * np.sqrt(periods_per_year))

def _calmar_ratio(cagr_percent: float, max_dd_percent: float) -> float:
    if max_dd_percent == 0:
        return 0.0
    return float(cagr_percent / abs(max_dd_percent))

def _trade_stats(trades_df):
    """Return winrate(%), avg win/loss, profit factor, trade count."""
    import pandas as pd

    if not isinstance(trades_df, pd.DataFrame):
        print("⚠️ Warning: trades_df is not a DataFrame. Skipping trade stats.")
        return 0.0, 0.0, 0.0, 0.0, 0

    if trades_df.empty or "Profit" not in trades_df.columns:
        print("⚠️ No valid trades or 'Profit' column missing.")
        return 0.0, 0.0, 0.0, 0.0, 0

    profits = pd.to_numeric(trades_df["Profit"], errors="coerce")
    valid = profits.dropna()
    if valid.empty:
        print("⚠️ Profit column has no realized values after coercion.")
        return 0.0, 0.0, 0.0, 0.0, 0

    wins = valid[valid > 0]
    losses = valid[valid < 0]

    winrate = 100.0 * len(wins) / len(valid)
    avg_win = wins.mean() if not wins.empty else 0.0
    avg_loss = losses.mean() if not losses.empty else 0.0
    gross_profit = wins.sum()
    gross_loss = abs(losses.sum())
    profit_factor = gross_profit / gross_loss if gross_loss > 0 else np.nan
    num_trades = len(valid)

    return winrate, avg_win, avg_loss, profit_factor, num_trades


def format_confusion_matrix_string(value):
    """Normalize confusion-matrix text for wrapped display."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""

    text = str(value).strip()
    if not text:
        return ""

    digits = re.findall(r"-?\d+", text)
    if len(digits) == 9:
        numbers = [int(d) for d in digits]
        rows = [numbers[i:i + 3] for i in range(0, 9, 3)]
    else:
        normalized = re.sub(r"[\[\];|]", " ", text.replace("\\n", "\n"))
        tokens = [int(tok) for tok in re.findall(r"-?\d+", normalized)]
        if len(tokens) % 3 != 0 or not tokens:
            return text  # fallback to original representation
        rows = [tokens[i:i + 3] for i in range(0, len(tokens), 3)]

    col_widths = []
    if rows:
        num_cols = len(rows[0])
        for col_idx in range(num_cols):
            col_widths.append(max(len(str(row[col_idx])) for row in rows))

    def _format_row(row, indent=False):
        pieces = []
        for col_idx, val in enumerate(row):
            width = col_widths[col_idx] if col_widths else len(str(val))
            part = str(val).rjust(width)
            if col_idx > 0 and not part.startswith(" "):
                part = " " + part
            pieces.append(part)
        line = " ".join(pieces)
        return (" " + line) if indent else line

    formatted_lines = []
    for idx, row in enumerate(rows):
        formatted_lines.append(_format_row(row, indent=(idx > 0)))

    if not formatted_lines:
        return text

    formatted_lines[0] = "[" + formatted_lines[0]
    formatted_lines[-1] = formatted_lines[-1] + "]"
    return "\n".join(formatted_lines)


def _max_line_length(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return 0
    string_value = str(value)
    if not string_value:
        return 0
    return max(len(line) for line in string_value.splitlines())


def apply_summary_excel_formatting(ws, confusion_col_name="Confusion_Matrix"):
    """Apply column widths, wrapping, alignment, and borders."""
    thin_side = Side(style="thin", color="D0D0D0")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    header_map = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        header_map[cell.value] = col_idx
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 32
    confusion_col_idx = header_map.get(confusion_col_name)

    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 50 if confusion_col_idx else 28
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = thin_border
            if confusion_col_idx and col_idx == confusion_col_idx:
                cell.font = Font(name="Consolas")

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row in range(1, ws.max_row + 1):
            max_len = max(max_len, _max_line_length(ws.cell(row=row, column=col_idx).value))
        width = min(max_len + 4, 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    print("[DEBUG] Excel formatting applied.")


def _equity_from_prices(close: pd.Series) -> pd.Series:
    # Build a buy&hold equity curve starting at STARTING_BALANCE
    if close is None or len(close) == 0:
        return pd.Series([STARTING_BALANCE])
    base = float(close.iloc[0])
    eq = (close.astype(float) / base) * STARTING_BALANCE
    return eq

def _sharpe_from_equity(equity_curve: pd.Series, rf_annual: float = 0.0, periods_per_year: int = 252) -> float:
    eq = pd.Series(equity_curve).astype(float).reset_index(drop=True)
    rets = eq.pct_change().dropna()
    if len(rets) < 2:
        return 0.0
    rf = rf_annual / periods_per_year
    excess = rets - rf
    denom = excess.std(ddof=1)
    if denom == 0 or np.isnan(denom):
        return 0.0
    return float((excess.mean() / denom) * np.sqrt(periods_per_year))

# ---------------- Main Summary Writer ---------------- #
def create_and_save_backtest_summary_table_csv_file(
    summary_tickers,                      # list[str]
    summary_CAGRs,                        # list[float]   (final CAGR of strategy)
    summary_buy_hold_CAGRs,               # list[float]
    summary_cagr_strategy_efficiencies,   # list[float]
    summary_sharpe_ratios,                # list[float]   (strategy Sharpe)
    # new optional containers:
    strategy_names=None,                  # list[str]
    model_numbers=None,                   # list[str] or list[int]
    feature_set_ids=None,                 # list[str]
    plot_links=None,                      # list[str] (file paths)
    portfolio_values_dict=None,           # dict[ticker]-> pd.Series (strategy equity curve)
    trade_tables_dict=None,               # dict[ticker]-> pd.DataFrame (must include 'Profit')
    price_series_dict=None,               # dict[ticker]-> pd.Series of 'Close' prices (for B&H Sharpe)
    ml_metrics_dict=None                  # dict[(ticker,strategy)]-> {accuracy, precision, recall, f1, confusion_matrix}
):
    """
    NOTE: We keep the original function name so your imports in main.py don't break.
    Output is now an Excel file with ALL metrics and a clickable Plot link.
    """
    os.makedirs(FINAL_EXCEL_DIR, exist_ok=True)
    file_name = f"final_backtesting_summary_{TICKERS}.xlsx"
    full_path = os.path.join(FINAL_EXCEL_DIR, file_name)

    rows = []
    n = len(summary_tickers)

    for i in range(n):
        ticker = summary_tickers[i]
        cagr = float(summary_CAGRs[i]) if summary_CAGRs[i] is not None else 0.0
        bh_cagr = float(summary_buy_hold_CAGRs[i]) if summary_buy_hold_CAGRs[i] is not None else 0.0
        cagr_eff = float(summary_cagr_strategy_efficiencies[i]) if summary_cagr_strategy_efficiencies[i] is not None else 0.0
        sharpe = float(summary_sharpe_ratios[i]) if summary_sharpe_ratios[i] is not None else 0.0

        strategy = strategy_names[i] if strategy_names else ""
        model_no = model_numbers[i] if model_numbers else ""
        feat_id  = feature_set_ids[i] if feature_set_ids else ""
        plot     = plot_links[i] if plot_links else ""

        # Trading metrics computed here (no new imports in main.py)
        pv_series = portfolio_values_dict.get(ticker) if portfolio_values_dict else None
        pv_len = len(pv_series) if pv_series is not None else 0
        print(f"[DEBUG] {ticker} portfolio_values length: {pv_len}")
        trades_df = trade_tables_dict.get(ticker) if trade_tables_dict else None
        print(
            f"[DEBUG] {ticker} trades_df type/shape: {type(trades_df)} / "
            f"{getattr(trades_df, 'shape', None)}"
        )
        close_ser = price_series_dict.get(ticker) if price_series_dict else None

        # Buy & hold Sharpe (computed on a normalized equity curve)
        bh_equity = _equity_from_prices(close_ser) if close_ser is not None else pd.Series([STARTING_BALANCE])
        bh_sharpe = _sharpe_from_equity(bh_equity, rf_annual=RISK_FREE_RATE)
        bh_mdd = _max_drawdown(bh_equity)

        sharpe_eff = float(sharpe / bh_sharpe) if bh_sharpe not in (0.0, np.nan) else 0.0
        sortino = _sortino_ratio(pv_series, rf_annual=RISK_FREE_RATE)
        mdd = _max_drawdown(pv_series)
        calmar = _calmar_ratio(cagr, mdd)
        print(f"[DEBUG] {ticker} Max Drawdown computed: {mdd:.2f}%")
        if isinstance(trades_df, str):
            print(f"[DEBUG] {ticker} trades_df string value: {trades_df}")
        elif isinstance(trades_df, pd.DataFrame):
            print(f"[DEBUG] {ticker} trades_df columns: {list(trades_df.columns)}")
            print(trades_df.head(3))
            if "Profit" in trades_df.columns:
                print(f"[DEBUG] {ticker} trades_df['Profit'].describe():")
                print(trades_df["Profit"].describe())

        winrate, avg_win, avg_loss, profit_factor, num_trades = _trade_stats(trades_df)
        print(
            f"[DEBUG] {ticker} WinRate computed: {winrate:.2f}%, "
            f"avg_win: {avg_win:.2f}, avg_loss: {avg_loss:.2f}, "
            f"profit_factor: {profit_factor if not pd.isna(profit_factor) else 'nan'}, "
            f"num_trades: {num_trades}"
        )

        # ML metrics — filled ONLY if provided for (ticker,strategy)
        key = (ticker, strategy)
        ml = ml_metrics_dict.get(key, {}) if ml_metrics_dict else {}
        acc  = ml.get("accuracy", "")
        prec = ml.get("precision", "")
        rec  = ml.get("recall", "")
        f1   = ml.get("f1", "")
        conf = ml.get("confusion_matrix", "")

        rows.append({
            "Ticker": ticker,
            "Strategy": strategy,
            "Model_Number": model_no,
            "Feature_Set_ID": feat_id,
            "CAGR(%)": round(cagr, 2),
            "B&H CAGR(%)": round(bh_cagr, 2),
            "CAGR Efficiency": round(cagr_eff, 2),
            "Sharpe": round(sharpe, 2),
            "B&H Sharpe": round(bh_sharpe, 2),
            "Sharpe Efficiency": round(sharpe_eff, 2),
            "Sortino": round(sortino, 2),
            "Max_Drawdown(%)": round(mdd, 2),
            "B&H Max_Drawdown(%)": round(bh_mdd, 2),
            "Calmar": round(calmar, 2),
            "WinRate(%)": round(winrate, 2),
            "Avg_Win": round(avg_win, 2),
            "Avg_Loss": round(avg_loss, 2),
            "Profit_Factor": round(profit_factor, 2),
            "Number_of_Trades": int(num_trades),
            "Accuracy": acc,
            "Precision": prec,
            "Recall": rec,
            "F1_Score": f1,
            "Confusion_Matrix": conf,
            "Plot_Link": plot
        })

    df = pd.DataFrame(rows)
    csv_df = df.copy()
    if "Confusion_Matrix" in df.columns:
        df["Confusion_Matrix"] = df["Confusion_Matrix"].apply(format_confusion_matrix_string)
    print(f"[DEBUG] Summary columns: {list(df.columns)}")
    os.makedirs(FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE, exist_ok=True)
    csv_path = os.path.join(
        FOLDER_PATH_FOR_SUMMARY_BACKTEST_TABLE, f"summary_backtest_table_{TICKERS}.csv"
    )
    csv_df.to_csv(csv_path, index=False)
    print(f"✅ Summary CSV with metrics saved at: {csv_path}")

    # Make Plot_Link clickable
    def _xl_hyperlink(path: str):
        if not path:
            return ""
        # Excel-friendly absolute path hyperlink
        return f'=HYPERLINK("{path}", "Open Plot")'

    if "Plot_Link" in df.columns:
        df["Plot_Link"] = df["Plot_Link"].apply(_xl_hyperlink)

    # Save to Excel
    with pd.ExcelWriter(full_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Summary")
        ws = writer.sheets["Summary"]
        # bold header
        for cell in ws[1]:
            cell.font = Font(bold=True)
        apply_summary_excel_formatting(ws)

    # small CLI preview (optional)
    try:
        from tabulate import tabulate as _tab
        print("\n\033[1mSummary of all results:\033[0m")
        print(_tab(df.head(20), headers="keys", tablefmt="grid", floatfmt=".2f"))
    except:
        pass

    print(f"\n✅ Final backtesting summary saved at:\n{full_path}\n")
    # keep return signature identical to your previous function (table_str, path)
    return df.to_string(index=False), full_path
